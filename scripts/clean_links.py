"""清洗“外链合集”Excel，生成内置站点库种子数据与私有提交记录。

用法：
    python scripts/clean_links.py <xlsx 路径> [--out data]

输入约定（与现有表格一致）：
- 第 1 个工作表：未清洗的站点总集（A 列分类或站点名，B 列 URL，其余列为备注）。
- 表头含“链接 / 状态”的工作表：某个产品的真实提交记录，表名即产品名
  （列：链接、状态、备注、排名；“排名”为 TRUE 表示该站支持打榜）。
- 表头含“反链”的工作表：每三列一组（收录、反链、徽章 HTML），组内第三列表头为产品名。

输出：
- <out>/seed/sites.csv          站点库种子（可提交到仓库；不含任何原始备注）
- <out>/private/submissions.csv 各产品提交记录（含备注，疑似密码已脱敏；勿提交）
- <out>/private/badges.csv      各产品的反链徽章代码与放置情况（勿提交）
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

import openpyxl

# ---------- URL 规范化 ----------

TRACKING_PARAMS = re.compile(r"^(ref|utm_\w+|source|via|aff|affiliate)$", re.I)
URL_IN_TEXT = re.compile(r"https?:/{1,2}[^\s\"'<>]+", re.I)


def normalize_url(raw: str | None) -> str | None:
    """修复常见录入错误并去掉追踪参数，返回规范 URL；无法识别时返回 None。"""
    if not raw or not isinstance(raw, str):
        return None
    text = raw.strip()
    m = URL_IN_TEXT.search(text.replace(" ", ""))
    if not m:
        return None
    url = re.sub(r"^(https?):/{1,2}", r"\1://", m.group(0), flags=re.I)
    parts = urlsplit(url)
    host = parts.hostname
    if not host or "." not in host:
        return None
    query = urlencode([(k, v) for k, v in parse_qsl(parts.query) if not TRACKING_PARAMS.match(k)])
    path = parts.path.rstrip("/") or ""
    return urlunsplit(("https", host.lower(), path, query, ""))


def domain_of(url: str | None) -> str | None:
    if not url:
        return None
    host = urlsplit(url).hostname or ""
    return host[4:] if host.startswith("www.") else host or None


# ---------- 备注解析 ----------

PASSWORD_AFTER_KEYWORD = re.compile(r"((?:密码|password|pwd|pass)\s*[:：]?\s*)([!-~]{4,})", re.I)
ASCII_RUN = re.compile(r"[!-~]{8,}")


def _looks_like_secret(token: str) -> bool:
    if token.lower().startswith("http") or "@" in token and "." in token.split("@")[-1]:
        return False  # URL 或邮箱地址
    classes = sum(
        bool(re.search(p, token)) for p in (r"[a-z]", r"[A-Z]", r"\d", r"[^A-Za-z0-9]")
    )
    return classes >= 3


def redact(note: str) -> tuple[str, bool]:
    """隐藏疑似密码。返回 (脱敏文本, 是否发生脱敏)。"""
    hit = False

    def _kw(m: re.Match[str]) -> str:
        nonlocal hit
        hit = True
        return m.group(1) + "[已隐藏]"

    text = PASSWORD_AFTER_KEYWORD.sub(_kw, note)

    def _run(m: re.Match[str]) -> str:
        nonlocal hit
        if _looks_like_secret(m.group(0)):
            hit = True
            return "[已隐藏]"
        return m.group(0)

    text = ASCII_RUN.sub(_run, text)
    return text, hit


REVIEW_TIME = re.compile(
    r"(\d+\s*[-~～]\s*\d+\s*(?:天|周|个月|月|小时|days?|weeks?|months?)|"
    r"\d+\s*(?:天|周|个月|月|小时|days?|weeks?|months?)\s*(?:\+|以上|内)?)"
    r"(?!\s*解锁)",
    re.I,
)
ANY_DATE = re.compile(
    r"(20\d{2}\s*年?\s*)?\d{1,2}\s*月\s*\d{1,2}\s*日?|(?<!\d)\d{1,2}\.\d{1,2}(\.\d{1,2})?(?!\d)"
)
FULL_DATE = [
    re.compile(r"(20\d{2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})"),
    re.compile(r"(20\d{2})\s+(\d{1,2})\s*月\s*(\d{1,2})"),
    re.compile(r"(?<!\d)(\d{2})\.(\d{1,2})\.(\d{1,2})(?!\d)"),  # 26.6.22
]


@dataclass
class NoteFacts:
    login_method: str = ""
    pricing: str = ""
    requires_backlink: bool = False
    requires_verification: bool = False
    review_time_hint: str = ""
    scheduled_date: str = ""
    problem: str = ""


def parse_note(note: str) -> NoteFacts:
    f = NoteFacts()
    if not note:
        return f
    low = note.lower()
    if "谷歌" in note or "google" in low:
        f.login_method = "google"
    elif "验证码登录" in note:
        f.login_method = "email_code"
    elif "邮箱" in note or "密码" in note:
        f.login_method = "email_password"
    if "反向链接免费" in note or "反链" in note or "徽章" in note:
        f.requires_backlink = True
        f.pricing = "free_with_backlink"
    if "付费" in note:
        f.pricing = "paid_with_free_window" if "免" in note else "paid"
    if "dns" in low or "代码验证" in note:
        f.requires_verification = True
    if m := REVIEW_TIME.search(ANY_DATE.sub(" ", note)):
        f.review_time_hint = re.sub(r"\s+", "", m.group(0))
    for i, pat in enumerate(FULL_DATE):
        if m := pat.search(note):
            y, mo, d = (int(x) for x in m.groups())
            if i == 2:
                y += 2000
            if 1 <= mo <= 12 and 1 <= d <= 31:
                f.scheduled_date = f"{y:04d}-{mo:02d}-{d:02d}"
            break
    for kw in ("禁用", "无法提交", "服务器错误", "出错"):
        if kw in note:
            f.problem = kw
            break
    return f


# ---------- 读取工作表 ----------

STATUS_MAP = {"已收录": "listed", "审核中": "in_review", "被拒": "rejected", "失败": "failed"}
GENERIC_CATEGORIES = {"saas": "SaaS", "ai": "AI"}
NEEDS_REVIEW_LABELS = {"extra list (to be processed)"}
# 头部平台：流量和曝光集中、审核严格或需要经营发布时机，只生成提交材料由用户亲自提交，不走自动化。
HEAD_PLATFORMS = {
    "producthunt.com",
    "betalist.com",
    "alternativeto.net",
    "indiehackers.com",
    "crunchbase.com",
    "f6s.com",
    "g2.com",
    "capterra.com",
    "appsumo.com",
    "toolify.ai",
    "futurepedia.io",
    "theresanaiforthat.com",
    "futuretools.io",
    "saashub.com",
    "peerlist.io",
    "devhunt.org",
    "microlaunch.net",
    "stackshare.io",
    "github.com",
    "news.ycombinator.com",
    "reddit.com",
}


@dataclass
class Site:
    domain: str
    name: str = ""
    submit_url: str = ""
    category: str = ""
    pricing: str = ""
    login_method: str = ""
    requires_backlink: bool = False
    requires_verification: bool = False
    supports_launch: bool = False
    review_time_hint: str = ""
    problems: set[str] = field(default_factory=set)
    needs_review: bool = False
    sources: set[str] = field(default_factory=set)
    listed: int = 0
    in_review: int = 0
    submitted: int = 0

    def merge_facts(self, f: NoteFacts) -> None:
        self.login_method = self.login_method or f.login_method
        if f.pricing and (not self.pricing or f.pricing.startswith("paid")):
            self.pricing = f.pricing
        self.requires_backlink |= f.requires_backlink
        self.requires_verification |= f.requires_verification
        self.review_time_hint = self.review_time_hint or f.review_time_hint
        if f.problem:
            self.problems.add(f.problem)


def cell_text(v: object) -> str:
    return "" if v is None else str(v).strip()


def is_true(v: object) -> bool:
    return v is True or cell_text(v).upper() in {"TRUE", "是", "Y", "YES", "1"}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--out", type=Path, default=Path("data"))
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    sites: dict[str, Site] = {}
    stats = defaultdict(int)
    invalid_urls: list[str] = []

    def site(domain: str) -> Site:
        return sites.setdefault(domain, Site(domain=domain))

    # 1) 总集
    master = wb.worksheets[0]
    for row in master.iter_rows(min_row=1, values_only=True):
        if not any(cell_text(c) for c in row):
            continue
        label, raw_url = cell_text(row[0]), row[1] if len(row) > 1 else None
        url = normalize_url(cell_text(raw_url))
        if not url:
            if cell_text(raw_url):
                invalid_urls.append(cell_text(raw_url))
            continue
        if cell_text(raw_url) and not cell_text(raw_url).lower().startswith(("http://", "https://")):
            stats["master_fixed_url"] += 1
        stats["master_rows"] += 1
        s = site(domain_of(url))
        s.sources.add("master")
        s.submit_url = s.submit_url or url
        key = label.lower()
        if key in GENERIC_CATEGORIES:
            s.category = s.category or GENERIC_CATEGORIES[key]
        elif key in NEEDS_REVIEW_LABELS:
            s.needs_review = True
        elif label:
            s.name = s.name or label
        for extra in row[2:]:
            s.merge_facts(parse_note(cell_text(extra)))

    # 2) 各产品提交记录
    submissions: list[dict[str, str]] = []
    badge_ws = None
    for ws in wb.worksheets[1:]:
        header = [cell_text(c) for c in next(ws.iter_rows(max_row=1, values_only=True), ())]
        if "反链" in header:
            badge_ws = ws
            continue
        if not ("链接" in header and "状态" in header):
            continue
        idx = {h: i for i, h in enumerate(header) if h}
        product = ws.title.strip()
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw_url = cell_text(row[idx["链接"]])
            url = normalize_url(raw_url)
            if not url:
                if raw_url:
                    invalid_urls.append(raw_url)
                continue
            domain = domain_of(url)
            status = STATUS_MAP.get(cell_text(row[idx["状态"]]), "unknown")
            raw_note = cell_text(row[idx["备注"]]) if "备注" in idx else ""
            launch = is_true(row[idx["排名"]]) if "排名" in idx and idx["排名"] < len(row) else False
            facts = parse_note(raw_note)
            note, hit = redact(raw_note)
            stats["redacted_notes"] += hit

            s = site(domain)
            s.sources.add(product)
            s.merge_facts(facts)
            s.supports_launch |= launch
            s.submitted += 1
            s.listed += status == "listed"
            s.in_review += status == "in_review"

            submissions.append(
                {
                    "product": product,
                    "domain": domain,
                    "recorded_url": url,
                    "status": status,
                    "supports_launch": str(launch).lower(),
                    "scheduled_date": facts.scheduled_date,
                    "review_time_hint": facts.review_time_hint,
                    "note": note,
                    "note_had_secret": str(hit).lower(),
                }
            )

    # 3) 反链徽章
    badges: list[dict[str, str]] = []
    if badge_ws is not None:
        rows = list(badge_ws.iter_rows(values_only=True))
        header = [cell_text(c) for c in rows[0]]
        product_cols = [i for i in range(2, len(header), 3) if header[i] and header[i] not in {"收录", "反链"}]
        for row in rows[1:]:
            for pc in product_cols:
                html = cell_text(row[pc]) if pc < len(row) else ""
                if not html:
                    continue
                domain = domain_of(normalize_url(html))
                if not domain:
                    continue
                site(domain).requires_backlink = True
                site(domain).sources.add("badges")
                badges.append(
                    {
                        "product": header[pc],
                        "domain": domain,
                        "listed": str(is_true(row[pc - 2])).lower(),
                        "badge_placed": str(is_true(row[pc - 1])).lower(),
                        "badge_html": re.sub(r"\s+", " ", html).strip(),
                    }
                )

    # 4) 汇总站点
    for s in sites.values():
        if s.requires_backlink and not s.pricing:
            s.pricing = "free_with_backlink"
        if s.problems:
            s.needs_review = True
        if not s.submit_url:
            s.submit_url = f"https://{s.domain}"

    seed_dir, private_dir = args.out / "seed", args.out / "private"
    seed_dir.mkdir(parents=True, exist_ok=True)
    private_dir.mkdir(parents=True, exist_ok=True)

    site_fields = [
        "domain", "name", "submit_url", "category", "pricing", "login_method",
        "requires_backlink", "requires_verification", "supports_launch", "review_time_hint",
        "submit_mode", "status", "problems", "sources", "submitted", "listed", "in_review",
    ]
    with open(seed_dir / "sites.csv", "w", newline="", encoding="utf-8-sig") as fp:
        w = csv.DictWriter(fp, fieldnames=site_fields)
        w.writeheader()
        for s in sorted(sites.values(), key=lambda x: (-x.listed, -x.submitted, x.domain)):
            w.writerow(
                {
                    "domain": s.domain,
                    "name": s.name,
                    "submit_url": s.submit_url,
                    "category": s.category,
                    "pricing": s.pricing or "unknown",
                    "login_method": s.login_method or "unknown",
                    "requires_backlink": str(s.requires_backlink).lower(),
                    "requires_verification": str(s.requires_verification).lower(),
                    "supports_launch": str(s.supports_launch).lower(),
                    "review_time_hint": s.review_time_hint,
                    "submit_mode": "manual" if s.domain in HEAD_PLATFORMS else "auto",
                    "status": "needs_review" if s.needs_review else "active",
                    "problems": "|".join(sorted(s.problems)),
                    "sources": "|".join(sorted(s.sources)),
                    "submitted": s.submitted,
                    "listed": s.listed,
                    "in_review": s.in_review,
                }
            )

    for name, rows_out in (("submissions.csv", submissions), ("badges.csv", badges)):
        if rows_out:
            with open(private_dir / name, "w", newline="", encoding="utf-8-sig") as fp:
                w = csv.DictWriter(fp, fieldnames=list(rows_out[0]))
                w.writeheader()
                w.writerows(rows_out)

    # 5) 报告
    by_status = defaultdict(int)
    for r in submissions:
        by_status[(r["product"], r["status"])] += 1
    products = sorted({r["product"] for r in submissions})
    print(f"站点（去重后）: {len(sites)}")
    print(f"  来自总集: {sum('master' in s.sources for s in sites.values())}")
    print(f"  仅出现在提交记录/徽章表: {sum('master' not in s.sources for s in sites.values())}")
    print(f"  有收录记录: {sum(s.listed > 0 for s in sites.values())}")
    print(f"  需复核: {sum(s.needs_review for s in sites.values())}")
    print(f"  支持打榜: {sum(s.supports_launch for s in sites.values())}")
    print(f"  需反链/徽章: {sum(s.requires_backlink for s in sites.values())}")
    print(f"  付费: {sum(s.pricing.startswith('paid') for s in sites.values())}")
    print(f"  需 Google 登录: {sum(s.login_method == 'google' for s in sites.values())}")
    print(f"总集中修复的 URL: {stats['master_fixed_url']}，无法识别: {len(invalid_urls)} {invalid_urls}")
    print(f"提交记录: {len(submissions)} 条，备注脱敏: {stats['redacted_notes']} 条")
    for p in products:
        parts = ", ".join(f"{k}={by_status[(p, k)]}" for k in ("listed", "in_review", "unknown") if by_status[(p, k)])
        print(f"  {p}: {parts}")
    print(f"徽章记录: {len(badges)} 条")
    print(f"输出: {seed_dir / 'sites.csv'}, {private_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
